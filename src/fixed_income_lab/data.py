from __future__ import annotations

import hashlib
import warnings
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook

from .paths import PROJECT_ROOT, raw_data_dir


TENOR_TO_YEARS = {
    "3月": 0.25,
    "6月": 0.5,
    "1年": 1.0,
    "2年": 2.0,
    "3年": 3.0,
    "5年": 5.0,
    "7年": 7.0,
    "10年": 10.0,
    "30年": 30.0,
}


class DataSourceError(RuntimeError):
    """Raised when an upstream market-data source cannot be normalized."""


@dataclass(frozen=True)
class MarketDataset:
    curve: pd.DataFrame
    issuance: pd.DataFrame
    news: pd.DataFrame


CHINABOND_CURVE_NAME = "中债国债收益率曲线（到期）"
CHINABOND_SOURCE = "ChinaBond/CCDC official annual file"
CHINABOND_ANNUAL_URL = (
    "https://yield.chinabond.com.cn/cbweb-mn/yc/downYearBzqx"
    "?locale=en_US&wrjxCBFlag=0"
    "&ycDefId=2c9081e50a2f9606010a3068cae70001"
    "&year={year}&zblx=txy"
)


def sample_dataset(data_dir: Path | None = None) -> MarketDataset:
    base = data_dir or PROJECT_ROOT / "data"
    return MarketDataset(
        curve=pd.read_csv(base / "sample_yield_curve.csv"),
        issuance=pd.read_csv(base / "sample_issuance.csv"),
        news=pd.read_csv(base / "sample_news.csv"),
    )


def _download_chinabond_annual_file(year: int, destination: Path) -> None:
    url = CHINABOND_ANNUAL_URL.format(year=year)
    request = Request(
        url,
        headers={
            "User-Agent": (
                "china-fixed-income-lab/0.1 "
                "(+https://github.com/SuperrrrSpike/china-fixed-income-lab)"
            )
        },
    )
    try:
        with urlopen(request, timeout=30) as response:
            content = response.read()
    except Exception as exc:
        raise DataSourceError(f"ChinaBond official download failed for {year}: {exc}") from exc

    if not content.startswith(b"PK"):
        raise DataSourceError(
            f"ChinaBond official download for {year} is not an OOXML workbook"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.name}.part")
    temporary.write_bytes(content)
    temporary.replace(destination)


def _find_column(frame: pd.DataFrame, candidates: tuple[str, ...], label: str) -> str:
    match = next((column for column in candidates if column in frame.columns), None)
    if match is None:
        raise DataSourceError(
            f"ChinaBond workbook is missing {label}. Received: {list(frame.columns)}"
        )
    return match


def parse_official_chinabond_workbook(
    workbook_path: Path,
    source_url: str,
) -> pd.DataFrame:
    """Parse a ChinaBond annual workbook into an audited long-form curve."""
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style",
                category=UserWarning,
            )
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook.active
        # ChinaBond files declare A1 as their used range even though they contain
        # thousands of rows. Recalculate the range before streaming the cells.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise DataSourceError("ChinaBond workbook contains no header row")
        raw = pd.DataFrame(rows, columns=header)
    except DataSourceError:
        raise
    except Exception as exc:
        raise DataSourceError(f"ChinaBond workbook could not be parsed: {exc}") from exc
    finally:
        if "workbook" in locals():
            workbook.close()

    date_column = _find_column(raw, ("Date", "日期"), "date column")
    tenor_column = _find_column(
        raw,
        ("Standard Terms(Yrs)", "标准期限(年)", "标准期限（年）"),
        "standard-tenor column",
    )
    yield_column = _find_column(raw, ("Yield(%)", "收益率(%)", "收益率（%）"), "yield column")

    result = pd.DataFrame(
        {
            "date": pd.to_datetime(raw[date_column], errors="coerce"),
            "tenor_years": pd.to_numeric(raw[tenor_column], errors="coerce"),
            "yield_pct": pd.to_numeric(raw[yield_column], errors="coerce"),
        }
    ).dropna()
    if result.empty:
        raise DataSourceError("ChinaBond workbook contains no valid curve observations")
    if result.duplicated(["date", "tenor_years"]).any():
        raise DataSourceError("ChinaBond workbook contains duplicate date-tenor observations")
    if not result["yield_pct"].between(-5.0, 30.0).all():
        raise DataSourceError("ChinaBond workbook contains implausible yield values")
    if not result["tenor_years"].between(0.0, 100.0).all():
        raise DataSourceError("ChinaBond workbook contains implausible tenor values")

    retrieved_at = datetime.fromtimestamp(
        workbook_path.stat().st_mtime,
        tz=timezone.utc,
    ).isoformat()
    source_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    result["curve_name"] = CHINABOND_CURVE_NAME
    result["source"] = CHINABOND_SOURCE
    result["source_url"] = source_url
    result["source_sha256"] = source_sha256
    result["retrieved_at_utc"] = retrieved_at
    result["data_status"] = "official_primary"
    columns = [
        "curve_name",
        "date",
        "tenor_years",
        "yield_pct",
        "source",
        "source_url",
        "source_sha256",
        "retrieved_at_utc",
        "data_status",
    ]
    return result[columns].sort_values(["date", "tenor_years"]).reset_index(drop=True)


def fetch_official_chinabond_curve(
    start_date: str,
    end_date: str,
    cache_dir: Path | None = None,
    force_refresh: bool = False,
) -> pd.DataFrame:
    """Download official ChinaBond annual files and return the requested range."""
    start = pd.to_datetime(start_date, errors="raise")
    end = pd.to_datetime(end_date, errors="raise")
    if start > end:
        raise ValueError("start_date must not be after end_date")

    cache = cache_dir or raw_data_dir()
    frames = []
    for year in range(start.year, end.year + 1):
        source_url = CHINABOND_ANNUAL_URL.format(year=year)
        workbook_path = cache / f"chinabond_government_yield_curve_{year}.xlsx"
        if force_refresh or not workbook_path.exists():
            _download_chinabond_annual_file(year, workbook_path)
        frames.append(parse_official_chinabond_workbook(workbook_path, source_url))

    result = pd.concat(frames, ignore_index=True)
    result = result[result["date"].between(start.normalize(), end.normalize())]
    if result.empty:
        raise DataSourceError(
            f"ChinaBond official files contain no rows from {start.date()} to {end.date()}"
        )
    return result.reset_index(drop=True)


def fetch_china_yield_curve(
    start_date: str,
    end_date: str,
    curve_keyword: str = "国债收益率曲线",
) -> pd.DataFrame:
    """Fetch a ChinaBond curve through AKShare as a non-authoritative fallback."""
    try:
        import akshare as ak
    except ImportError as exc:
        raise DataSourceError("AKShare is not installed. Run pip install -e '.[dev]'.") from exc

    try:
        frame = ak.bond_china_yield(start_date=start_date, end_date=end_date)
    except Exception as exc:  # Network and upstream errors need contextual wrapping.
        raise DataSourceError(f"AKShare yield-curve request failed: {exc}") from exc

    if frame.empty:
        raise DataSourceError("AKShare returned no yield-curve rows")

    curve_column = next(
        (name for name in ("曲线名称", "curve_name") if name in frame.columns),
        None,
    )
    if curve_column is None:
        raise DataSourceError(f"Curve-name column missing. Received: {list(frame.columns)}")

    matches = frame[curve_column].astype(str).str.contains(curve_keyword, na=False)
    selected = frame.loc[matches].copy()
    if selected.empty:
        available = frame[curve_column].dropna().astype(str).unique()[:10]
        raise DataSourceError(
            f"No curve matched {curve_keyword!r}. Available examples: {available.tolist()}"
        )

    selected = selected.rename(columns={curve_column: "curve_name", "日期": "date"})
    selected["source"] = "AKShare proxy for ChinaBond"
    tenor_columns = [column for column in TENOR_TO_YEARS if column in selected.columns]
    if len(tenor_columns) < 4:
        raise DataSourceError(
            f"Too few supported tenor columns. Received: {list(selected.columns)}"
        )
    expected = ["curve_name", "date", *tenor_columns, "source"]
    return selected[expected].sort_values("date").reset_index(drop=True)
